from __future__ import annotations

import hashlib
import json
import math
import re
from dataclasses import dataclass
from functools import lru_cache
from io import BytesIO
from typing import Any, Mapping
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import coordinate_to_tuple

from .additional_options import ADDITIONAL_OPTIONS


SCHEMA_VERSION = 1
MAX_WORKBOOK_BYTES = 20 * 1024 * 1024
MAX_HISTORY_ROWS = 20_000
NUMERIC_FIELDS = {
    "capacity": "capacity_kg", "speed": "speed_ms", "stops": "stops",
    "rise_mm": "lifting_height_mm", "landing_doors": "doors_count",
    "cabin_width": "cabin_width_mm", "cabin_depth": "cabin_depth_mm",
    "cabin_height": "cabin_height_mm", "door_width": "landing_door_width_mm",
    "door_height": "landing_door_height_mm", "shaft_width": "shaft_width_mm",
    "shaft_depth": "shaft_depth_mm", "pit": "pit_depth_mm", "overhead": "overhead_mm",
}
CATEGORY_FIELDS = {
    "machine_room": "machine_room", "door_type": "door_opening_type",
    "wall_side": "side_wall_finish", "wall_rear": "rear_wall_finish",
    "finish_other": "other_floors_landing_door_finish", "firefighters": "firefighter_mode",
    "fire_resistance": "fire_resistance", "mgn": "mgn_accessibility",
    "display": "display_type", "group_control": "group_operation", "door_brand": "door_model",
    "cop": "cop_type", "ceiling": "ceiling_type", "floor": "floor_finish",
    "mirror": "mirror", "cwt_safety": "room_under_pit", "seismic": "seismic",
}
OPTION_FIELDS = {
    "ard": "option_ard", "access": "option_ic_card", "cat_cable": "option_cctv_preparation",
    "ado": "option_ado", "noise": "option_extra_noise_insulation",
    "remote_control": "option_remote_control_cabinet", "logo": "option_customer_logo",
    "independent": "option_extra_lop_line", "video": "option_video_camera",
}
PRICING_GROUP_FIELDS = frozenset(
    [*NUMERIC_FIELDS.values(), *CATEGORY_FIELDS.values(), *OPTION_FIELDS.values(),
     "cabin_type", "main_floor_landing_door_finish", "additional_options"]
)
REQUIRED_INPUT_LABELS = {
    "capacity": "грузоподъёмность", "speed": "скорость", "stops": "остановки",
    "rise_mm": "высота подъёма", "cabin_width": "ширина кабины",
    "cabin_depth": "глубина кабины", "cabin_height": "высота кабины",
    "door_width": "ширина двери", "door_height": "высота двери",
}


class PriceCatalogError(ValueError):
    pass


@dataclass(frozen=True)
class PriceFeature:
    key: str
    cell: str
    weight: float
    scale: float | None = None


@dataclass(frozen=True)
class PriceRecord:
    group: int
    price: float
    period: float | None
    video: bool
    values: tuple[float | str | None, ...]


@dataclass(frozen=True)
class PriceEstimate:
    cny: int
    raw_cny: float
    nearest_distance: float


@dataclass(frozen=True, eq=False)
class PriceCatalog:
    source_name: str
    source_date: str
    source_sha256: str
    latest_period: float
    time_coefficient: float
    weight_decay: float
    weight_offset: float
    weight_power: float
    video_package_cny: float
    features: tuple[PriceFeature, ...]
    records: tuple[PriceRecord, ...]

    def inputs_for_group(self, group: Mapping[str, Any]) -> tuple[Any, ...]:
        inputs: dict[str, Any] = {key: _positive_number(group.get(field)) for key, field in NUMERIC_FIELDS.items()}
        inputs.update({key: _category(key, group.get(field)) for key, field in CATEGORY_FIELDS.items()})
        inputs["entrances"] = 2.0 if _text(group.get("cabin_type")) == "проходная" else 1.0
        inputs["quantity"] = 1.0
        inputs["regeneration"] = "Нет"
        inputs["compact_control"] = "Нет"
        if inputs["landing_doors"] is None:
            inputs["landing_doors"] = inputs["stops"]
        if not _text(group.get("other_floors_landing_door_finish")):
            inputs["finish_other"] = _material(group.get("main_floor_landing_door_finish"))
        selected_lines = set(str(group.get("additional_options") or "").splitlines())
        for key, field in OPTION_FIELDS.items():
            if field in group:
                selected = _yes(group[field])
            else:
                selected = any(
                    option.field == field and (option.chinese in selected_lines or option.russian in selected_lines)
                    for option in ADDITIONAL_OPTIONS
                )
            inputs[key] = "Да" if selected else "Нет"
        return tuple(inputs.get(feature.key, "Не указано") for feature in self.features) + (inputs["video"],)

    def missing_fields(self, inputs: tuple[Any, ...]) -> tuple[str, ...]:
        values = {feature.key: value for feature, value in zip(self.features, inputs)}
        return tuple(label for key, label in REQUIRED_INPUT_LABELS.items() if _positive_number(values.get(key)) is None)

    def estimate_group(self, group: Mapping[str, Any]) -> PriceEstimate | None:
        return self.estimate_inputs(self.inputs_for_group(group))

    @lru_cache(maxsize=512)
    def estimate_inputs(self, inputs: tuple[Any, ...]) -> PriceEstimate | None:
        if len(inputs) != len(self.features) + 1 or self.missing_fields(inputs):
            return None
        best_by_group: dict[int, tuple[float, int, PriceRecord]] = {}
        for index, record in enumerate(self.records, start=2):
            distance = 0.0
            for feature, actual, wanted in zip(self.features, record.values, inputs):
                if feature.scale is None:
                    delta = float(str(actual or "").casefold() != str(wanted or "").casefold())
                elif actual is None or wanted is None:
                    delta = 1.0
                else:
                    delta = min(2.0, abs(float(actual) - float(wanted)) / feature.scale)
                distance += feature.weight * delta
            candidate = (distance, index, record)
            previous = best_by_group.get(record.group)
            if previous is None or candidate[:2] < previous[:2]:
                best_by_group[record.group] = candidate
        # Excel keeps one record per independent group and uses the source row to break ties.
        nearest = sorted(best_by_group.values(), key=lambda item: item[0] + item[1] / 100_000_000)[:5]
        weights = [math.exp(-self.weight_decay * distance) / (self.weight_offset + distance) ** self.weight_power
                   for distance, _, _ in nearest]
        weight_sum = sum(weights)
        adjusted_prices = [
            record.price * math.exp(0 if record.period is None else
                                   (self.latest_period - record.period) * self.time_coefficient)
            for _, _, record in nearest
        ]
        average = sum(weight * price for weight, price in zip(weights, adjusted_prices)) / weight_sum
        video_share = sum(weight for weight, (_, _, record) in zip(weights, nearest) if record.video) / weight_sum
        price = max(0.0, average - self.video_package_cny * video_share)
        if inputs[-1] == "Да":
            price += self.video_package_cny
        return PriceEstimate(int(math.floor(price / 100 + 0.5)) * 100, price, nearest[0][0])

    def to_bytes(self) -> bytes:
        payload = {
            "schema_version": SCHEMA_VERSION,
            "source_name": self.source_name, "source_date": self.source_date,
            "source_sha256": self.source_sha256, "latest_period": self.latest_period,
            "time_coefficient": self.time_coefficient, "weight_decay": self.weight_decay,
            "weight_offset": self.weight_offset, "weight_power": self.weight_power,
            "video_package_cny": self.video_package_cny,
            "features": [[f.key, f.cell, f.weight, f.scale] for f in self.features],
            "records": [[r.group, r.price, r.period, r.video, list(r.values)] for r in self.records],
        }
        return json.dumps(payload, ensure_ascii=False, separators=(",", ":"), allow_nan=False).encode("utf-8")

    @classmethod
    def from_bytes(cls, content: bytes) -> PriceCatalog:
        try:
            data = json.loads(content)
            if data["schema_version"] != SCHEMA_VERSION:
                raise ValueError("schema")
            catalog = cls(
                source_name=str(data["source_name"]), source_date=str(data["source_date"]),
                source_sha256=str(data["source_sha256"]), latest_period=_finite(data["latest_period"]),
                time_coefficient=_finite(data["time_coefficient"]), weight_decay=_finite(data["weight_decay"]),
                weight_offset=_finite(data["weight_offset"]), weight_power=_finite(data["weight_power"]),
                video_package_cny=_finite(data["video_package_cny"]),
                features=tuple(PriceFeature(str(k), str(c), _finite(w), None if s is None else _finite(s))
                               for k, c, w, s in data["features"]),
                records=tuple(PriceRecord(int(g), _finite(p), None if t is None else _finite(t), bool(v), tuple(values))
                              for g, p, t, v, values in data["records"]),
            )
            catalog._validate()
            return catalog
        except (ValueError, TypeError, KeyError, OverflowError) as error:
            raise PriceCatalogError("Файл прайса повреждён или имеет неподдерживаемый формат.") from error

    def _validate(self) -> None:
        if not 5 <= len(self.records) <= MAX_HISTORY_ROWS or len({r.group for r in self.records}) < 5:
            raise PriceCatalogError("В прайсе нужно не менее пяти независимых групп с ценами.")
        if not 1 <= len(self.features) <= 60 or len({f.key for f in self.features}) != len(self.features):
            raise PriceCatalogError("Некорректные параметры модели цен.")
        if not set(REQUIRED_INPUT_LABELS) <= {f.key for f in self.features}:
            raise PriceCatalogError("В модели отсутствуют обязательные параметры лифта.")
        supported = set(NUMERIC_FIELDS) | set(CATEGORY_FIELDS) | set(OPTION_FIELDS) | {
            "entrances", "quantity", "regeneration", "compact_control",
        }
        if any(feature.key not in supported for feature in self.features):
            raise PriceCatalogError("В модели появился новый параметр цены. Нужна адаптация импорта.")
        if self.weight_offset <= 0 or self.weight_decay < 0 or self.weight_power <= 0 or self.video_package_cny < 0:
            raise PriceCatalogError("Некорректные коэффициенты модели цен.")
        if not 1900 <= self.latest_period <= 2200 or abs(self.time_coefficient) > 2:
            raise PriceCatalogError("Некорректная временная поправка цен.")
        for feature in self.features:
            if feature.weight <= 0 or (feature.scale is not None and feature.scale <= 0):
                raise PriceCatalogError("Некорректный вес параметра цены.")
        for record in self.records:
            if record.price <= 0 or len(record.values) != len(self.features):
                raise PriceCatalogError("Некорректная историческая цена.")
            if record.period is not None and not 1900 <= record.period <= 2200:
                raise PriceCatalogError("Некорректная дата исторической цены.")
            for feature, value in zip(self.features, record.values):
                if feature.scale is not None and value is not None:
                    _finite(value)
                elif feature.scale is None and not isinstance(value, str):
                    raise PriceCatalogError("Некорректная категория исторической конфигурации.")


def catalog_from_configurator(content: bytes, source_name: str) -> PriceCatalog:
    if len(content) > MAX_WORKBOOK_BYTES:
        raise PriceCatalogError("Конфигуратор должен быть не больше 20 МБ.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 120 * 1024 * 1024:
                raise PriceCatalogError("Слишком большой объём данных в конфигураторе.")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
        try:
            return _catalog_from_workbook(workbook, content, source_name)
        finally:
            workbook.close()
    except PriceCatalogError:
        raise
    except (BadZipFile, KeyError, ValueError, TypeError, IndexError, OverflowError, StopIteration) as error:
        raise PriceCatalogError("Нужен файл EPSS_PRICE_CONFIGURATOR.xlsx с расчётной базой цен.") from error


def _catalog_from_workbook(workbook, content: bytes, source_name: str) -> PriceCatalog:
    config = list(workbook["КОНФИГУРАТОР"].values)
    calc = list(workbook["_CALC"].values)
    near = list(workbook["_NEAR"].values)
    headers = near[0]
    checks = {
        "D6": '=IF(_CALC!B20,ROUND(_CALC!R10/100,0)*100,"Заполните обязательные поля")',
        "_CALC!R10": "=R9+R8", "_CALC!R9": "=MAX(0,B9-R7)",
        "_CALC!R8": '=IF(КОНФИГУРАТОР!B41="Да",R5,0)',
        "_CALC!R7": "=R5*R6", "_CALC!R6": "=SUMPRODUCT(M2:M6,S2:S6)/SUM(M2:M6)",
        "_CALC!R5": "=КОНФИГУРАТОР!B48", "_CALC!B9": "=SUMPRODUCT(M2:M6,N2:N6)/SUM(M2:M6)",
    }
    for address, expected in checks.items():
        if "!" in address:
            _, cell = address.split("!")
            row, col = coordinate_to_tuple(cell)
            actual = calc[row - 1][col - 1]
        else:
            row, col = coordinate_to_tuple(address)
            actual = config[row - 1][col - 1]
        if _formula_signature(actual) != _formula_signature(expected):
            raise PriceCatalogError("В конфигураторе изменился способ расчёта цены. Нужна адаптация импорта.")
    terms = _formula_terms(near[1][7])
    features = tuple(_feature_from_term(term, headers) for term in terms)
    feature_columns = [headers.index(feature.key) for feature in features]
    ranking_ranges = [token.value for token in Tokenizer(calc[1][10]).items if token.subtype == "RANGE"]
    history_range = next(value for value in ranking_ranges if "_NEAR" in value and ":" in value)
    history_end = coordinate_to_tuple(history_range.split(":")[-1].replace("$", ""))[0]
    if history_end != len(near):
        raise PriceCatalogError("Диапазон расчёта не охватывает всю базу конфигуратора.")
    analog_checks = (
        (calc[1][10], f"=MATCH(SMALL(_NEAR!I2:I{history_end},J2),_NEAR!I2:I{history_end},0)"),
        (calc[1][11], f"=INDEX(_NEAR!H2:H{history_end},K2)"),
        (calc[1][13], f"=INDEX(_NEAR!K2:K{history_end},K2)"),
        (calc[1][18], f'=IF(INDEX(_NEAR!BG2:BG{history_end},K2)="Да",1,0)'),
        (near[1][8], f'=IF(COUNTIFS(B2:B{history_end},B2,H2:H{history_end},"<"&H2)'
         '+COUNTIFS(B2:B2,B2,H2:H2,H2)=1,H2+ROW()/100000000,1000000+ROW())'),
        (near[1][10], '=C2*EXP(IF(D2="",0,(_CALC!B2-D2)*_CALC!B3))'),
    )
    for actual, expected in analog_checks:
        if _formula_signature(actual) != _formula_signature(expected):
            raise PriceCatalogError("В конфигураторе изменился подбор аналогов. Нужна адаптация импорта.")
    weight_tokens = Tokenizer(calc[1][12]).items
    weight_numbers = [token.value for token in weight_tokens if token.subtype == "NUMBER"]
    decay, offset, power = map(_finite, weight_numbers)
    expected_weight = f"=EXP(-{weight_numbers[0]}*L2)/({weight_numbers[1]}+L2)^{weight_numbers[2]}"
    if _formula_signature(calc[1][12]) != _formula_signature(expected_weight):
        raise PriceCatalogError("Неподдерживаемая формула весов аналогов.")
    group_ids = {
        key: index for index, key in enumerate(dict.fromkeys(str(row[1]).casefold() for row in near[1:]))
    }
    records = tuple(
        PriceRecord(
            group_ids[str(row[1]).casefold()], _finite(row[2]), None if row[3] in (None, "") else _finite(row[3]),
            _yes(row[headers.index("Видеонаблюдение")]),
            tuple(None if row[column] in (None, "") else
                  _finite(row[column]) if feature.scale is not None else str(row[column])
                  for feature, column in zip(features, feature_columns)),
        )
        for row in near[1:]
    )
    date_label = str(config[1][3] or "")
    source_date = date_label.rsplit(":", 1)[-1].strip()
    catalog = PriceCatalog(
        source_name=source_name, source_date=source_date,
        source_sha256=hashlib.sha256(content).hexdigest(),
        latest_period=_finite(calc[1][1]), time_coefficient=_finite(calc[2][1]),
        weight_decay=decay, weight_offset=offset, weight_power=power,
        video_package_cny=_finite(config[47][1]), features=features, records=records,
    )
    catalog._validate()
    # Verify the imported calculation against the saved Excel example before accepting an update.
    raw_inputs = tuple(
        config[coordinate_to_tuple(feature.cell)[0] - 1][1]
        if not feature.cell.startswith("_CALC!") else calc[coordinate_to_tuple(feature.cell.split("!")[1])[0] - 1][17]
        for feature in features
    ) + (config[40][1],)
    cached = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        expected_price = cached["КОНФИГУРАТОР"]["D6"].value
        expected_raw_price = cached["_CALC"]["R10"].value
    finally:
        cached.close()
    estimate = catalog.estimate_inputs(raw_inputs)
    if isinstance(expected_price, (int, float)) and (estimate is None or estimate.cny != expected_price):
        raise PriceCatalogError("Расчёт не совпал с сохранённой ценой Excel. Пересчитайте и сохраните конфигуратор.")
    if isinstance(expected_raw_price, (int, float)) and (
        estimate is None or not math.isclose(estimate.raw_cny, expected_raw_price, rel_tol=1e-9, abs_tol=0.01)
    ):
        raise PriceCatalogError("Расчёт не совпал с сохранённой ценой Excel. Пересчитайте и сохраните конфигуратор.")
    return catalog


def _formula_signature(formula: Any) -> tuple[tuple[str, str, Any], ...]:
    if not isinstance(formula, str) or not formula.startswith("="):
        raise PriceCatalogError("Отсутствуют формулы расчёта цен.")
    signature = []
    for token in Tokenizer(formula).items:
        if token.type == "WHITE-SPACE":
            continue
        value: Any = token.value
        if token.subtype == "NUMBER":
            value = _finite(value)
        elif token.subtype == "RANGE":
            value = value.replace("$", "").replace("'", "").upper()
        signature.append((token.type, token.subtype, value))
    return tuple(signature)


def _formula_terms(formula: str) -> list[list[Any]]:
    terms: list[list[Any]] = [[]]
    depth = 0
    for token in Tokenizer(formula).items:
        if token.subtype == "OPEN":
            depth += 1
        elif token.subtype == "CLOSE":
            depth -= 1
        if token.value == "+" and depth == 0:
            terms.append([])
        elif token.type != "WHITE-SPACE":
            terms[-1].append(token)
    return terms


def _feature_from_term(tokens: list[Any], headers: tuple[Any, ...]) -> PriceFeature:
    weight = _finite(tokens[0].value)
    references = [token.value for token in tokens if token.subtype == "RANGE"]
    local, input_ref = references[:2]
    row, column = coordinate_to_tuple(local.replace("$", ""))
    key = headers[column - 1]
    numeric = any(token.value == "MIN(" for token in tokens)
    scale_text = next((token.value for token in reversed(tokens) if token.subtype == "NUMBER"), None) if numeric else None
    if numeric:
        expected = f'={tokens[0].value}*IF(OR({local}="",{input_ref}=""),1,MIN(2,ABS({local}-{input_ref})/{scale_text}))'
    else:
        expected = f'={tokens[0].value}*IF({local}={input_ref},0,1)'
    actual = "=" + "".join(token.value for token in tokens)
    if row != 2 or _formula_signature(actual) != _formula_signature(expected):
        raise PriceCatalogError("Неподдерживаемая формула сходства конфигураций.")
    sheet, coordinate = input_ref.replace("$", "").replace("'", "").split("!")
    if sheet == "КОНФИГУРАТОР":
        cell = coordinate
    elif sheet == "_CALC" and coordinate in {"R1", "R2"}:
        cell = f"_CALC!{coordinate}"
    else:
        raise PriceCatalogError("Неподдерживаемый параметр расчёта цены.")
    return PriceFeature(str(key), cell, weight, None if scale_text is None else _finite(scale_text))


def _finite(value: Any) -> float:
    number = float(value)
    if not math.isfinite(number):
        raise PriceCatalogError("В прайсе обнаружено некорректное число.")
    return number


def _positive_number(value: Any) -> float | None:
    try:
        number = float(str(value).strip().replace(",", "."))
        return number if math.isfinite(number) and number > 0 else None
    except (ValueError, TypeError):
        return None


def _text(value: Any) -> str:
    return str(value or "").strip().casefold().replace("ё", "е")


def _yes(value: Any) -> bool:
    return value is True or _text(value) in {"да", "yes", "true", "1", "есть"}


def _material(value: Any) -> str:
    text = _text(value)
    if not text:
        return "Не указано"
    if any(word in text for word in ("под отделку", "заказчик", "local", "customer")):
        return "Под отделку заказчика"
    if any(word in text for word in ("стекл", "glass", "панорам")) or ("зеркал" in text and "сталь" not in text):
        return "Стекло / зеркало"
    if "нержав" in text or "stainless" in text or "ex-hs" in text or "ex-ms" in text:
        return "Нержавеющая сталь"
    if any(word in text for word in ("крашен", "окраш", "paint")) or "ex-ys" in text:
        return "Крашеная сталь"
    return "Другая отделка"


def _category(key: str, value: Any) -> str:
    text = _text(value)
    if key in {"wall_side", "wall_rear", "finish_other"}:
        return _material(value)
    if key in {"firefighters", "mgn", "cwt_safety", "seismic"}:
        if not text:
            return "Нет"
        return "Да" if _yes(value) or (key == "seismic" and text not in {"нет", "0", "не указано"}) else "Нет"
    if not text:
        return "Не указано"
    if key == "machine_room":
        return "Без МП (MRL)" if "без" in text or "mrl" in text else "С МП (MR)"
    if key == "door_type":
        for fragment, category in (("телескоп", "Телескопическое"), ("централ", "Центральное"),
                                   ("вертикал", "Вертикальное"), ("распаш", "Распашное")):
            if fragment in text:
                return category
        return "Другое открывание"
    if key == "group_control":
        return "DDS" if "dds" in text else "Групповая" if "групп" in text else "Одиночная" if "одиноч" in text else "Другой режим"
    if key == "fire_resistance":
        if text in {"нет", "без огнестойкости"}:
            return "Без огнестойкости"
        match = re.fullmatch(r"(ei|e)[\s-]*(30|60|90)", text)
        return f"{match[1].upper()} {match[2]}" if match else "Другой класс"
    if key == "display":
        return "ЖК-дисплей (LCD / TFT)" if "lcd" in text or "tft" in text or "жк" in text else "LED" if "led" in text else "Без дисплея" if text == "нет" else "Другой дисплей"
    if key == "door_brand":
        for brand in ("NBSL", "Fermator", "Wittur"):
            if brand.casefold() in text:
                return brand
        return "Другой производитель"
    if key == "cop":
        return "Серия AC99" if "ac99" in text else "Серия AC55 / AC56" if "ac55" in text or "ac56" in text else "Под отделку заказчика" if "под отделку" in text else "Другая панель"
    if key == "ceiling":
        return "Под отделку заказчика" if "под отделку" in text else "Заводской потолок" if "ex-" in text else "Другой потолок"
    if key == "floor":
        if "под отделку" in text:
            return "Под отделку заказчика"
        if any(word in text for word in ("pvc", "пвх", "резин", "прорезин")):
            return "ПВХ / резина"
        if any(word in text for word in ("камень", "керам", "мрамор", "гранит", "stone")):
            return "Камень / керамогранит"
        return "Металл" if "сталь" in text or "металл" in text else "Другое покрытие"
    if key == "mirror":
        return "Нет" if "без" in text or text == "нет" else "Да"
    return str(value).strip()
