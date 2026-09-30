from dataclasses import replace
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

import app
from src.pricing import (
    REQUIRED_INPUT_LABELS, PriceCatalog, PriceCatalogError, PriceFeature, PriceRecord,
    catalog_from_configurator,
)


@pytest.fixture(scope="module")
def catalog():
    return PriceCatalog.from_bytes((Path(__file__).resolve().parents[1] / "data/epss_price_catalog.json").read_bytes())


@pytest.fixture
def group():
    return {
        "quantity": 2, "capacity_kg": 1000, "speed_ms": "1,6", "stops": 7,
        "lifting_height_mm": 18900, "doors_count": 7, "cabin_type": "Непроходная",
        "cabin_width_mm": 1100, "cabin_depth_mm": 1400, "cabin_height_mm": 2400,
        "landing_door_width_mm": 900, "landing_door_height_mm": 2200,
        "shaft_width_mm": 1700, "shaft_depth_mm": 1850, "pit_depth_mm": 1200, "overhead_mm": 3600,
        "machine_room": "Без машинного помещения", "door_opening_type": "Телескопическое",
        "side_wall_finish": "Шлифованная нержавеющая сталь EX-HS01",
        "rear_wall_finish": "Шлифованная нержавеющая сталь EX-HS01",
        "other_floors_landing_door_finish": "Шлифованная нержавеющая сталь EX-HS01",
        "firefighter_mode": "НЕТ", "fire_resistance": "EI-60", "mgn_accessibility": "НЕТ",
        "display_type": 'LCD 15"', "group_operation": "Групповая", "door_model": "NBSL",
        "cop_type": "EX-AC55", "ceiling_type": "EX-J135", "floor_finish": "Прорезиненное покрытие, PVC EX-DB210",
        "mirror": "MEX-3", "room_under_pit": "НЕТ", "seismic": "НЕТ",
        "option_cctv_preparation": "ДА", "option_ado": "ДА", "option_video_camera": "ДА",
    }


def test_saved_excel_control_price_matches_to_the_cent(catalog, group):
    estimate = catalog.estimate_group(group)
    assert estimate.cny == 91700
    assert estimate.raw_cny == pytest.approx(91749.21252447162, abs=0.01)


def test_quantity_does_not_change_unit_price_and_video_adds_package(catalog, group):
    estimate = catalog.estimate_group(group)
    assert catalog.estimate_group({**group, "quantity": 20}) == estimate
    without_video = catalog.estimate_group({**group, "option_video_camera": "НЕТ"})
    assert estimate.raw_cny - without_video.raw_cny == pytest.approx(catalog.video_package_cny)


@pytest.mark.parametrize("field,value", [
    ("capacity_kg", None), ("speed_ms", ""), ("stops", 0),
    ("lifting_height_mm", "nan"), ("cabin_width_mm", "РАСЧЁТНОЕ"),
    ("cabin_depth_mm", "МИНИМАЛЬНОЕ"), ("landing_door_height_mm", "МАКСИМАЛЬНОЕ"),
])
def test_incomplete_or_text_dimensions_do_not_get_invented_prices(catalog, group, field, value):
    inputs = catalog.inputs_for_group({**group, field: value})
    assert catalog.estimate_inputs(inputs) is None
    assert catalog.missing_fields(inputs)


def test_materials_and_legacy_option_lines_are_normalized(catalog, group):
    inputs = catalog.inputs_for_group(group)
    values = dict(zip((f.key for f in catalog.features), inputs))
    assert values["fire_resistance"] == "EI 60"
    assert values["speed"] == 1.6
    assert values["wall_side"] == "Нержавеющая сталь"
    assert values["display"] == "ЖК-дисплей (LCD / TFT)"
    assert values["cat_cable"] == "Да"
    legacy = dict(group)
    del legacy["option_video_camera"]
    legacy["additional_options"] = "视频摄像头"
    assert catalog.inputs_for_group(legacy) == inputs
    assert catalog.inputs_for_group({**group, "side_wall_finish": group["side_wall_finish"] + " AFP"}) == inputs


def test_independent_analogs_and_original_row_tiebreak():
    features = tuple(PriceFeature(key, "B6", 1.0, 100.0) for key in REQUIRED_INPUT_LABELS)
    values = (100.0,) * len(features)
    records = tuple(PriceRecord(i, 10000 + 1000 * i, None, False, values) for i in range(5))
    catalog = PriceCatalog("test.xlsx", "30.09.2026", "", 2026, 0, 3, 0.03, 2, 3500, features, records)
    catalog._validate()
    inputs = values + ("Нет",)
    assert catalog.estimate_inputs(inputs).cny == 12000
    duplicates = tuple(PriceRecord(0, 99000, None, True, values) for _ in range(8))
    assert replace(catalog, records=records + duplicates).estimate_inputs(inputs).cny == 12000
    better = replace(records[0], values=(101.0,) + values[1:])
    modified = replace(catalog, records=(better,) + records[1:] + (records[0],))
    assert modified.estimate_inputs(inputs).cny == 12000


def test_catalog_roundtrip_is_identical_and_has_no_source_project_names(catalog):
    restored = PriceCatalog.from_bytes(catalog.to_bytes())
    assert restored.records == catalog.records
    assert restored.features == catalog.features
    assert restored.source_sha256 == catalog.source_sha256
    assert len(catalog.records) == 1544
    assert b'"project"' not in catalog.to_bytes()
    assert b'"customer"' not in catalog.to_bytes()


@pytest.mark.parametrize("content", [b"not json", b"{}", b'{"schema_version":99}'])
def test_invalid_catalog_is_reported(content):
    with pytest.raises(PriceCatalogError):
        PriceCatalog.from_bytes(content)


def test_new_model_feature_requires_an_explicit_mapping(catalog):
    changed = replace(catalog, features=catalog.features + (PriceFeature("new_unmapped_option", "B50", 1.0),))
    with pytest.raises(PriceCatalogError, match="новый параметр"):
        changed._validate()


def test_summary_separates_different_pricing_configurations(catalog, group):
    rows = app._lift_summary_rows([
        group, {**group, "quantity": 3, "section": "B"},
        {**group, "quantity": 1, "option_video_camera": "НЕТ"},
    ], catalog)
    assert len(rows) == 2
    assert rows[0]["label"].startswith("5 лифтов")
    assert rows[1]["label"].startswith("1 лифт")
    assert rows[0]["cny"] - rows[1]["cny"] == 3500


def test_summary_reads_live_fragment_fields_without_resetting_widgets(monkeypatch, catalog, group):
    class State(dict):
        def __getattr__(self, name):
            return self[name]

        def __setattr__(self, name, value):
            self[name] = value

    state = State({
        "group_count": 1, "prefill_groups": [group], "group_drafts": [{}],
        "group_0_option_video_camera": False, "group_0_cabin_width_mm": "1150",
    })
    monkeypatch.setattr(app.st, "session_state", state)
    monkeypatch.setattr(app, "_price_catalog", lambda: catalog)
    summary = app._project_summary_from_state()
    expected = catalog.estimate_group({**group, "option_video_camera": False, "cabin_width_mm": "1150"})
    assert summary["lift_prices"][0]["cny"] == expected.cny
    assert state["group_0_option_video_camera"] is False
    assert state["group_0_cabin_width_mm"] == "1150"


def test_sidebar_formats_unit_prices_in_supplied_placeholder(monkeypatch):
    rendered = []

    class Placeholder:
        def markdown(self, content, **kwargs):
            rendered.append(content)

    monkeypatch.setattr(app, "_project_summary_from_state", lambda: {
        "project_name": "Объект", "lift_count": 2, "firefighter_lift_count": 0, "mgn_lift_count": 0,
        "lift_breakdown": ["2 лифта — 1,6 м/с, 1000 кг, 7 ост."],
        "lift_prices": [{"cny": 63000, "note": 'Предварительная заводская цена за один лифт. Прайс от "30.09.2026".'}],
    })
    app._render_project_summary_sidebar(Placeholder())
    assert len(rendered) == 1
    assert "~ 63 000 ¥" in rendered[0]
    assert "Предварительно за 1 лифт" not in rendered[0]
    assert 'title="Предварительная заводская цена за один лифт.' in rendered[0]
    assert "&quot;30.09.2026&quot;" in rendered[0]


def _configurator_fixture():
    workbook = Workbook()
    config = workbook.active
    config.title = "КОНФИГУРАТОР"
    calc = workbook.create_sheet("_CALC")
    near = workbook.create_sheet("_NEAR")
    fields = tuple(REQUIRED_INPUT_LABELS)
    input_rows = (6, 7, 8, 9, 12, 13, 14, 15, 16)
    for offset, (field, row) in enumerate(zip(fields, input_rows), start=16):
        config.cell(row, 2, 100)
        near.cell(1, offset, field)
        for record_row in range(2, 7):
            near.cell(record_row, offset, 100)
    near.cell(1, 59, "Видеонаблюдение")
    for row in range(2, 7):
        near.cell(row, 2, f"CUSTOMER_OBJECT_{row}")
        near.cell(row, 3, 10000 + row * 100)
        near.cell(row, 4, 2026)
        near.cell(row, 59, "Нет")
    near["H2"] = "=" + "+".join(
        f'1*IF(OR({get_column_letter(col)}2="",КОНФИГУРАТОР!B{row}=""),1,'
        f'MIN(2,ABS({get_column_letter(col)}2-КОНФИГУРАТОР!B{row})/100))'
        for col, row in enumerate(input_rows, start=16)
    )
    config["D2"] = "Последняя загрузка цен: 30.09.2026"
    config["D6"] = '=IF(_CALC!B20,ROUND(_CALC!R10/100,0)*100,"Заполните обязательные поля")'
    config["B41"] = "Нет"
    config["B48"] = 3500
    calc["B2"] = 2026
    calc["B3"] = 0
    formulas = {
        "R10": "=R9+R8", "R9": "=MAX(0,B9-R7)", "R8": '=IF(КОНФИГУРАТОР!B41="Да",R5,0)',
        "R7": "=R5*R6", "R6": "=SUMPRODUCT(M2:M6,S2:S6)/SUM(M2:M6)",
        "R5": "=КОНФИГУРАТОР!B48", "B9": "=SUMPRODUCT(M2:M6,N2:N6)/SUM(M2:M6)",
        "K2": "=MATCH(SMALL(_NEAR!I2:I6,J2),_NEAR!I2:I6,0)",
        "L2": "=INDEX(_NEAR!H2:H6,K2)", "M2": "=EXP(-3*L2)/(0.03+L2)^2",
        "N2": "=INDEX(_NEAR!K2:K6,K2)", "S2": '=IF(INDEX(_NEAR!BG2:BG6,K2)="Да",1,0)',
    }
    for address, formula in formulas.items():
        calc[address] = formula
    near["I2"] = '=IF(COUNTIFS(B2:B6,B2,H2:H6,"<"&H2)+COUNTIFS(B2:B2,B2,H2:H2,H2)=1,H2+ROW()/100000000,1000000+ROW())'
    near["K2"] = '=C2*EXP(IF(D2="",0,(_CALC!B2-D2)*_CALC!B3))'
    return workbook


def _xlsx_bytes(workbook):
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_configurator_import_reads_model_and_removes_sensitive_identifiers():
    catalog = catalog_from_configurator(_xlsx_bytes(_configurator_fixture()), "test.xlsx")
    assert catalog.source_date == "30.09.2026"
    assert len(catalog.records) == 5
    assert catalog.records[0].group == 0
    assert catalog.estimate_inputs((100.0,) * len(catalog.features) + ("Нет",)).cny == 10400
    assert b"CUSTOMER_OBJECT" not in catalog.to_bytes()


@pytest.mark.parametrize("sheet,cell", [
    ("КОНФИГУРАТОР", "D6"), ("_CALC", "K2"), ("_CALC", "M2"), ("_NEAR", "I2"), ("_NEAR", "H2"),
])
def test_changed_configurator_formula_is_not_silently_imported(sheet, cell):
    workbook = _configurator_fixture()
    workbook[sheet][cell] = "=123"
    with pytest.raises(PriceCatalogError):
        catalog_from_configurator(_xlsx_bytes(workbook), "test.xlsx")
